"""상위기관(교육(지원)청) 이첩용 '특이민원 관련 지원 요청 보고서'(매뉴얼 55쪽 서식)와 '이첩 기준 체크리스트'(매뉴얼 12쪽)를 만든다.

- 서식 항목: 학교명 / 제기대상 분류 / 민원관련 학생(성명·학급정보) / 보호자(성명·연락처) / 민원 제기 및 처리내용 /
  이첩 사유(학교 자체 해결 노력 포함) / 교육(지원)청 요청사항.
- 체크리스트: 매뉴얼 12쪽 '지원 요청 및 이첩 기준(안)' 5개 중 어느 것에 해당하는지를 민원 내용·이력으로 자동 판단한다.
  판단은 후보 제시용이며, 학교에서 한 조치(퇴거명령, 분쟁조정 등)처럼 시스템이 모르는 사항은 '확인 필요'로 남긴다.
- 쪽 번호는 매뉴얼에 인쇄된 번호(12쪽=PDF 16쪽, 55쪽=PDF 59쪽)다.
"""
import re

CRITERIA = [   # 매뉴얼 12쪽 '지원 요청 및 이첩 기준(안)' 원문
    "합리적인 답변을 하고 3회 이상 동일･유사한 민원을 제기하여 종료하였음에도 지속적으로 처리를 요구(자체 종결처리 가능)하거나 위력을 행사하는 경우",
    "위법･부당한 행위로 학교장이 퇴거명령, 민원 제한 등의 조치를 했음에도 불응하고 학교 및 상급기관에 지속적으로 민원을 제기하는 경우",
    "위법 행위가 관련 법령에 따라 처벌이 가능할 정도로 중대 사안인 경우",
    "학교에서 분쟁조정, 중재, 법률지원 등을 받았음에도 자체적으로 해결이 불가한 경우",
    "기타 학교의 정상적인 교육활동에 심각한 지장을 초래하는 경우",
]
FORM_CATEGORIES = ["교무/학사", "학부모/교직원 소통", "기타", "생활지도/상담", "학교 운영"]    # 서식의 제기대상 분류(서식 순서)
# 담당 부서 -> 서식의 제기대상 분류. 서식은 5분류라 부서 체계와 1:1이 아니며 자동 제안일 뿐이다.
DEPT_TO_FORM = {"교무기획부": "교무/학사", "교수학습지원부": "교무/학사", "고교학점제운영부": "교무/학사",
                "학생생활안전부": "생활지도/상담", "창의진로상담부": "생활지도/상담", "보건실": "생활지도/상담",
                "1학년부": "학부모/교직원 소통", "2학년부": "학부모/교직원 소통", "3학년부": "학부모/교직원 소통",
                "학년부(학년 확인 필요)": "학부모/교직원 소통",
                "행정실": "학교 운영", "급식실": "학교 운영", "교육정보부": "학교 운영", "예체능교육부": "학교 운영",
                "자연과학교육부": "학교 운영", "인문사회교육부": "학교 운영"}
FORCE_KEYS = {"assault", "weapon", "threat", "damage", "stalking"}      # '위력을 행사하는 경우'에 걸릴 수 있는 유형
GRAVE_SEVERITY = 4                                                        # 처벌 가능할 정도의 중대 사안 후보로 보는 심각도
STATE_MARK = {"해당": "■", "해당 가능": "▲", "확인 필요": "□", "해당 없음": "－"}
INPUT = "(담당자 입력)"


def _matched(h):
    """감지된 표현(문자열 또는 목록) -> 보여 줄 문자열."""
    m = h.get("matched") or ""
    return (m if isinstance(m, str) else ", ".join(m))[:40] or "-"


def _hits(result):
    return [h for h in result.get("malicious", []) if h.get("severity", 0) >= 3]


def checklist(result):
    """5개 이첩 기준마다 {no, text, state, reasons[]}. state: 해당 / 해당 가능 / 확인 필요 / 해당 없음."""
    repeat = result.get("repeat_count", 1)
    hist = result.get("history") or {}
    prior = hist.get("same_student") or []
    closed_before = any(p.get("closed") for p in prior)
    hits = _hits(result)
    force = [h for h in hits if h["key"] in FORCE_KEYS]
    grave = [h for h in hits if h.get("severity", 0) >= GRAVE_SEVERITY]
    out = []

    # 1) 3회 이상 반복 후 종료했는데도 지속 / 위력 행사
    reasons, state = [], "해당 없음"
    if repeat >= 3 and closed_before:
        state = "해당"
        ids = ", ".join(f"#{p['id']}" for p in prior if p.get("closed")) or "이전 민원"
        reasons.append(f"같은 학생의 동일·유사 민원이 이번까지 {repeat}회째이고, 이전에 종결(거부·종료)한 민원({ids})이 있는데도 다시 제기됨")
    elif repeat >= 3:
        state = "해당 가능"
        reasons.append(f"동일·유사 민원이 이번까지 {repeat}회째입니다. 이전 민원에 합리적인 답변을 하고 종료했는지 확인이 필요합니다")
    if force:
        state = state if state == "해당" else "해당 가능"
        reasons.append("위력 행사 가능성: " + ", ".join(f"{h['label']}(감지 표현: {_matched(h)})" for h in force))
    out.append((state, reasons or ["3회 이상 반복·위력 행사에 해당하는 자동 감지 결과가 없습니다"]))

    # 2) 퇴거명령·민원 제한에도 불응하고 지속 제기 (학교가 한 조치는 시스템이 모른다)
    pm = hist.get("prior_malicious", 0)
    r2 = ["학교장의 퇴거명령·민원 제한 조치 여부와 그 이후의 불응·지속 제기는 시스템에 기록이 없어 담당자가 확인해야 합니다"]
    if pm:
        r2.append(f"참고: 같은 학생의 과거 악성(특이) 의심 민원이 {pm}건 있습니다")
    out.append(("확인 필요", r2))

    # 3) 처벌 가능한 중대 사안
    if grave:
        sure = all(h.get("confident", True) for h in grave)
        out.append(("해당" if sure else "해당 가능",
                    [f"{h['label']}(심각도 {h['severity']}): 감지 표현 '{_matched(h)}'" for h in grave]
                    + ["처벌 가능 여부와 적용 법령은 교육(지원)청·법률 검토로 확정해야 합니다"]))
    else:
        out.append(("해당 없음", ["처벌 가능한 중대 사안으로 볼 만한 유형(심각도 4 이상)이 감지되지 않았습니다"]))

    # 4) 분쟁조정·중재·법률지원을 받았는데도 해결 불가
    out.append(("확인 필요", ["학교가 분쟁조정·중재·법률지원을 받았는지는 시스템에 기록이 없어 담당자가 확인해야 합니다"]))

    # 5) 정상적인 교육활동에 심각한 지장
    r5, st5 = [], "확인 필요"
    if any(h["key"] == "stalking" for h in hits):
        st5 = "해당 가능"
        r5.append("스토킹 유형이 감지되어 교직원 업무·교육활동에 지장을 줄 가능성이 있습니다")
    if repeat >= 3:
        st5 = "해당 가능"
        r5.append(f"동일·유사 민원이 {repeat}회째 반복되어 담당 교직원의 업무에 지장을 줄 가능성이 있습니다")
    r5.append("실제로 교육활동에 심각한 지장이 있었는지는 담당자가 사실을 확인해야 합니다")
    out.append((st5, r5))
    return [{"no": i + 1, "text": CRITERIA[i], "state": s, "reasons": r} for i, (s, r) in enumerate(out)]


def _short(text, n=120):
    text = " ".join((text or "").split())
    return text[:n] + ("…" if len(text) > n else "")


def build_transfer(result, complaint=None, school_name=""):
    """보고서 내용. result는 분류 결과, complaint는 store.get_complaint 결과(있으면 처리 경과·접수일을 쓴다)."""
    cl = checklist(result)
    chosen = [c for c in cl if c["state"] in ("해당", "해당 가능")]
    dept = result.get("department", "")
    form_cat = DEPT_TO_FORM.get(dept, "기타")
    student = result.get("student")
    m = re.match(r"(.+)\((\d{3})\)$", student or "")
    name, cls = (m.group(1), m.group(2)) if m else ("", "")
    klass = f"{cls[0]}학년 {int(cls[1:])}반({cls})" if cls else ""
    hist = result.get("history") or {}
    created = (complaint or {}).get("created_at") or "접수 시각 미기록"

    # 민원 제기 및 처리내용: 과거 같은 학생의 유사 민원 -> 이번 민원 순
    lines = []
    for p in sorted(hist.get("same_student") or [], key=lambda p: p["id"]):
        outcome = "종결(거부·종료) 처리" if p.get("closed") else "접수·처리"
        lines.append(f"{p['created_at'][:10]} · 민원 #{p['id']}: {p['snippet']} → {outcome}")
    cid = result.get("complaint_id")
    treated = [e["note"] for e in (complaint or {}).get("events", []) if e.get("note")]
    lines.append(f"{created[:10]} · 이번 민원{f' #{cid}' if cid else ''}: {_short(result.get('text_masked'))} → "
                 + (" / ".join(treated[:3]) if treated else f"{result.get('route_label', '')}, 담당 {result.get('handler', '')}"))

    # 이첩 사유 (작성예시 문형을 따르고, 해당 기준 번호를 적는다)
    if chosen:
        nos = ", ".join(f"{c['no']}번" for c in chosen)
        why = [f"체크리스트(매뉴얼 12쪽 '지원 요청 및 이첩 기준(안)') {nos} 기준에 해당합니다."]
        why += [f"- {c['no']}번 [{c['state']}] {'; '.join(c['reasons'][:2])}" for c in chosen]
        if result.get("repeat_count", 1) >= 3:
            why.append(f"(작성예시 문형) 합리적인 답변을 했음에도 불구하고 3회 이상 동일·유사한 민원을 지속 제기하여 처리를 요구하거나 "
                       f"위력을 행사하였기에 이첩 및 지원 요청합니다 (동일·유사 민원 {result.get('repeat_count')}회째).")
    else:
        why = ["현재 자동 판단으로는 이첩 기준에 해당하는 항목이 없습니다. 이첩이 필요하다고 판단한 사유를 담당자가 직접 적어 주세요."]
    effort = [f"[{e['created_at']}] {e['kind']}{' · ' + e['status'] if e.get('status') else ''}{': ' + e['note'] if e.get('note') else ''}"
              for e in (complaint or {}).get("events", []) if e["kind"] != "접수"]
    efforts = effort or [INPUT + " (학교 자체 해결 노력: 답변·면담·경고문 발송 등)"]
    asks = []       # 교육(지원)청에 요청할 사항(초안): 선택된 기준과 대응 단계에서 끌어온다
    if 1 in [c["no"] for c in chosen]:
        asks.append("동일·유사 민원 반복 및 위력 행사 우려에 대한 이첩 접수와 향후 민원 처리 방향 협의")
    if 3 in [c["no"] for c in chosen]:
        asks.append("고소·고발 등 법적 조치 협의 및 법률 지원 (매뉴얼 50쪽 '학교장 대응' 단계: 교육(지원)청 관련 부서와 협의)")
    if 5 in [c["no"] for c in chosen]:
        asks.append("피해 교직원 보호조치와 교육활동 침해행위 해당 여부 검토 지원")
    if any("출입·면담 제한" in a["text"] for a in result.get("actions", [])):
        asks.append("출입·면담 제한 등 접촉 차단 방안 검토")
    requests = [f"- {x}" for x in asks] + ["(위 내용은 선택된 기준으로 만든 초안입니다. 학교 상황에 맞게 고쳐 주세요)"] if asks else [INPUT]
    return {"school": school_name or INPUT, "category": form_cat, "form_categories": FORM_CATEGORIES,
            "dept": dept, "student_name": name or INPUT, "class_info": klass or INPUT,
            "guardian_name": INPUT, "guardian_contact": INPUT, "history_lines": lines,
            "reason_lines": why, "effort_lines": efforts, "request_lines": requests,
            "checklist": cl, "chosen": [c["no"] for c in chosen]}


def to_docx(t, path):
    """build_transfer 결과를 매뉴얼 55쪽 서식 모양의 Word 파일로 쓴다."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "맑은 고딕", Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    doc.add_heading("특이민원 관련 지원 요청 보고서", level=1)
    p = doc.add_paragraph("※ 학교민원 처리 매뉴얼 55쪽 서식(교육(지원)청별 여건에 따라 변경 가능). 자동 분류로 만든 초안이며 학교장 확인 후 제출합니다.")
    p.runs[0].font.size = Pt(8.5)

    def row(table, label, lines):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = ""
        for i, ln in enumerate(lines if isinstance(lines, list) else [lines]):
            para = cells[1].paragraphs[0] if i == 0 else cells[1].add_paragraph()
            para.add_run(ln)
        for r in cells[0].paragraphs[0].runs:
            r.bold = True

    tb = doc.add_table(rows=0, cols=2)
    tb.style = "Table Grid"
    row(tb, "학교명", t["school"])
    row(tb, "제기대상 분류", "  ".join(f"{'☑' if c == t['category'] else '☐'} {c}" for c in t["form_categories"])
        + f"\n(자동 제안, 담당 부서: {t['dept']})")
    row(tb, "민원관련 학생", [f"성명: {t['student_name']}", f"학급정보: {t['class_info']}"])
    row(tb, "보호자", [f"성명: {t['guardian_name']}", f"연락처: {t['guardian_contact']}"])
    row(tb, "민원제기 및 처리내용", t["history_lines"])
    row(tb, "이첩 사유\n(학교 자체 해결\n노력 내용 포함)", t["reason_lines"] + ["", "학교 자체 해결 노력:"] + t["effort_lines"])
    row(tb, "교육(지원)청\n요청사항", t["request_lines"])
    for r in tb.rows:
        r.cells[0].width = r.cells[0].width // 2
    doc.add_paragraph("※ 이 서식 외에 NEIS '특이민원관리' 메뉴를 통한 이첩 및 지원 요청도 가능합니다(매뉴얼 55쪽).").runs[0].font.size = Pt(8.5)

    doc.add_page_break()
    doc.add_heading("별첨. 이첩 기준 체크리스트 (매뉴얼 12쪽 '지원 요청 및 이첩 기준(안)')", level=2)
    doc.add_paragraph("■ 해당   ▲ 해당 가능(사람이 확인)   □ 확인 필요(학교 기록 필요)   － 해당 없음")
    ct = doc.add_table(rows=1, cols=4)
    ct.style = "Table Grid"
    for c, h in zip(ct.rows[0].cells, ("번호", "이첩 기준(매뉴얼 원문)", "판단", "이첩 사유·근거")):
        c.text = h
        c.paragraphs[0].runs[0].bold = True
    for c in t["checklist"]:
        cells = ct.add_row().cells
        cells[0].text = str(c["no"])
        cells[1].text = c["text"]
        cells[2].text = f"{STATE_MARK[c['state']]} {c['state']}"
        cells[3].text = "\n".join(c["reasons"])
    doc.add_paragraph()
    sel = ", ".join(f"{n}번" for n in t["chosen"]) or "없음"
    doc.add_paragraph(f"선택된 기준(해당·해당 가능): {sel}")
    doc.save(path)
    return path
